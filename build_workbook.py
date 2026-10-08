"""Builds the Google-Sheets-ready workbook (Nova_Marketing_Governance_DB.xlsx).

Import it into Google Sheets (File > Import) to get the same database the prototype runs on. The metric columns of
CAMPAIGN_PERFORMANCE and the DASHBOARD_DATA KPIs are real spreadsheet FORMULAS (COUNTIFS / SUMIF / IF), exactly as
they would be in Google Sheets - nothing is pasted as a value.

Run directly:  python google_sheets/build_workbook.py   (writes google_sheets/Nova_Marketing_Governance_DB.xlsx)
"""
import io
import os
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from engine import constants as C  # noqa: E402

INR = '[>=10000000]"₹"##\\,##\\,##\\,##0;[>=100000]"₹"##\\,##\\,##0;"₹"##,##0'
NAVY = "1F2A44"


def _sheet(wb, name, headers, rows, widths=None, money=(), pct=(), first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = name
    ws.append(headers)
    for r in rows:
        ws.append(list(r))
    for c in range(1, len(headers) + 1):
        h = ws.cell(row=1, column=c)
        h.font = Font(bold=True, color="FFFFFF")
        h.fill = PatternFill("solid", fgColor=NAVY)
        h.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = (widths or {}).get(headers[c - 1], 16)
    ws.freeze_panes = "B2"
    ws.row_dimensions[1].height = 30
    for name_ in money:
        col = headers.index(name_) + 1
        for r in range(2, ws.max_row + 1):
            ws.cell(row=r, column=col).number_format = INR
    for name_ in pct:
        col = headers.index(name_) + 1
        for r in range(2, ws.max_row + 1):
            ws.cell(row=r, column=col).number_format = "0.0%"
    return ws


def build(conn):
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "README"
    notes = [
        "Nova Consumer Brands - Marketing Campaign Governance (Google Sheets database)",
        "",
        "SIMULATED academic data. Import this file into Google Sheets: File > Import > Upload > 'Replace spreadsheet'.",
        "Then open Extensions > Apps Script and paste the .gs files from /apps_script (see docs/GOOGLE_DEPLOYMENT.md).",
        "",
        "Sheets: CAMPAIGN_REQUESTS, APPROVAL_RULES, APPROVAL_LOG, LEAD_DATA, CAMPAIGN_PERFORMANCE (formulas), USERS, AUDIT_TRAIL,",
        "NOTIFICATIONS (outbox), SETTINGS, LISTS (dropdowns), DASHBOARD_DATA (KPI formulas = Looker Studio data source).",
        "",
        "Protected in production: APPROVAL_LOG, AUDIT_TRAIL (script-write only), Approval_Status / Execution_Status columns (script-write only).",
    ]
    for i, t in enumerate(notes, 1):
        ws0.cell(row=i, column=1, value=t)
    ws0["A1"].font = Font(bold=True, size=14)
    ws0.column_dimensions["A"].width = 130

    # ---- CAMPAIGN_REQUESTS
    cols = ["Campaign_ID", "Submission_Date", "Region", "Requester_Name", "Requester_Email", "Campaign_Name", "Product",
            "Campaign_Type", "Objective", "Target_Audience", "Marketing_Channel", "Start_Date", "End_Date", "Proposed_Budget",
            "Expected_Leads", "Expected_Conversions", "Expected_Revenue", "Campaign_Description", "Creative_Link",
            "Approval_Level", "Current_Approver", "Approval_Status", "Execution_Status", "Rejection_Reason",
            "Created_Timestamp", "Updated_Timestamp", "Approval_Chain", "Approval_Step", "Rule_ID", "Revision_No",
            "Submitted_Timestamp", "Approved_Timestamp", "Ready_Timestamp", "Executed_Timestamp"]
    rows = [[r[c] for c in cols] for r in conn.execute("SELECT * FROM campaign_requests ORDER BY Campaign_ID")]
    w = {"Campaign_Name": 34, "Target_Audience": 40, "Campaign_Description": 50, "Creative_Link": 36, "Rejection_Reason": 44,
         "Marketing_Channel": 34, "Current_Approver": 30, "Requester_Email": 30, "Approval_Chain": 40}
    ws = _sheet(wb, "CAMPAIGN_REQUESTS", cols, rows, w, money=("Proposed_Budget", "Expected_Revenue"))
    n = ws.max_row

    # ---- APPROVAL_RULES
    rcols = ["Rule_ID", "Rule_Name", "Min_Budget", "Min_Inclusive", "Max_Budget", "Max_Inclusive", "Approval_Chain", "Active"]
    _sheet(wb, "APPROVAL_RULES", rcols, [[r[c] for c in rcols] for r in conn.execute("SELECT * FROM approval_rules")],
           {"Rule_Name": 34, "Approval_Chain": 70}, money=("Min_Budget", "Max_Budget"))

    # ---- APPROVAL_LOG
    acols = ["Approval_ID", "Campaign_ID", "Approver_Name", "Approver_Role", "Approval_Level", "Decision", "Decision_Date",
             "Comments", "Previous_Status", "New_Status", "Revision_No"]
    _sheet(wb, "APPROVAL_LOG", acols, [[r[c] for c in acols] for r in conn.execute("SELECT * FROM approval_log ORDER BY Approval_ID")],
           {"Comments": 50, "Decision_Date": 20, "Approver_Role": 22})

    # ---- LEAD_DATA
    lcols = ["Lead_ID", "Campaign_ID", "Region", "Lead_Source", "Lead_Date", "Customer_Segment", "Lead_Score", "Lead_Status",
             "Sales_Assigned", "Conversion_Status", "Revenue"]
    ws_l = _sheet(wb, "LEAD_DATA", lcols, [[r[c] for c in lcols] for r in conn.execute("SELECT * FROM lead_data ORDER BY Lead_ID")],
                  {"Customer_Segment": 22, "Lead_Source": 20, "Sales_Assigned": 20}, money=("Revenue",))
    ln = ws_l.max_row

    # ---- SETTINGS
    srows = [[r["Key"], r["Value"], r["Description"]] for r in conn.execute("SELECT * FROM settings")]
    ws_s = _sheet(wb, "SETTINGS", ["Key", "Value", "Description"], srows, {"Key": 30, "Description": 90})
    th_row = 1 + [r[0] for r in srows].index("qualified_score_threshold") + 1

    # ---- CAMPAIGN_PERFORMANCE (inputs + FORMULAS)
    pcols = ["Campaign_ID", "Campaign_Name", "Region", "Budget", "Actual_Spend", "Impressions", "Clicks", "Leads",
             "Qualified_Leads", "Conversions", "Revenue", "CPL", "Conversion_Rate", "ROAS", "ROI"]
    ws_p = wb.create_sheet("CAMPAIGN_PERFORMANCE")
    ws_p.append(pcols)
    perf = conn.execute("SELECT * FROM campaign_performance ORDER BY Campaign_ID").fetchall()
    for i, p in enumerate(perf, start=2):
        ws_p.append([
            p["Campaign_ID"],
            f'=IFERROR(INDEX(CAMPAIGN_REQUESTS!$F$2:$F${n},MATCH(A{i},CAMPAIGN_REQUESTS!$A$2:$A${n},0)),"")',
            f'=IFERROR(INDEX(CAMPAIGN_REQUESTS!$C$2:$C${n},MATCH(A{i},CAMPAIGN_REQUESTS!$A$2:$A${n},0)),"")',
            f'=IFERROR(INDEX(CAMPAIGN_REQUESTS!$N$2:$N${n},MATCH(A{i},CAMPAIGN_REQUESTS!$A$2:$A${n},0)),0)',
            p["Actual_Spend"], p["Impressions"], p["Clicks"],
            f"=COUNTIF(LEAD_DATA!$B$2:$B${ln},A{i})",
            f'=COUNTIFS(LEAD_DATA!$B$2:$B${ln},A{i},LEAD_DATA!$G$2:$G${ln},">="&SETTINGS!$B${th_row})',
            f'=COUNTIFS(LEAD_DATA!$B$2:$B${ln},A{i},LEAD_DATA!$J$2:$J${ln},"Converted")',
            f"=SUMIF(LEAD_DATA!$B$2:$B${ln},A{i},LEAD_DATA!$K$2:$K${ln})",
            f'=IF(H{i}>0,E{i}/H{i},"")',          # CPL = Actual Spend / Leads
            f'=IF(H{i}>0,J{i}/H{i},"")',          # Conversion Rate = Conversions / Leads
            f'=IF(E{i}>0,K{i}/E{i},"")',          # ROAS = Revenue / Spend
            f'=IF(E{i}>0,(K{i}-E{i})/E{i},"")',   # ROI = (Revenue - Spend) / Spend
        ])
    for c in range(1, len(pcols) + 1):
        h = ws_p.cell(row=1, column=c)
        h.font = Font(bold=True, color="FFFFFF"); h.fill = PatternFill("solid", fgColor=NAVY)
        ws_p.column_dimensions[get_column_letter(c)].width = 34 if c == 2 else 16
    ws_p.freeze_panes = "B2"
    pn = ws_p.max_row
    for r in range(2, pn + 1):
        for c in (4, 5, 11):
            ws_p.cell(row=r, column=c).number_format = INR
        ws_p.cell(row=r, column=12).number_format = '"₹"#,##0.00'
        ws_p.cell(row=r, column=13).number_format = "0.0%"
        ws_p.cell(row=r, column=14).number_format = '0.00"x"'
        ws_p.cell(row=r, column=15).number_format = "0.0%"
        for c in (6, 7, 8, 9, 10):
            ws_p.cell(row=r, column=c).number_format = "#,##0"
    ws_p.conditional_formatting.add(f"O2:O{pn}", CellIsRule(operator="lessThan", formula=["0"], font=Font(color="C0392B", bold=True)))

    # ---- USERS / AUDIT / NOTIFICATIONS
    ucols = ["User_ID", "Name", "Email", "Role", "Region", "Active"]
    _sheet(wb, "USERS", ucols, [[r[c] for c in ucols] for r in conn.execute("SELECT * FROM users")], {"Name": 22, "Email": 32, "Role": 30})
    ucols = ["Audit_ID", "Timestamp", "User", "User_Role", "Action", "Campaign_ID", "Previous_Status", "New_Status", "Details"]
    _sheet(wb, "AUDIT_TRAIL", ucols, [[r[c] for c in ucols] for r in conn.execute("SELECT * FROM audit_trail ORDER BY Timestamp, Audit_ID")],
           {"Details": 80, "Timestamp": 20, "User": 20, "Action": 26, "User_Role": 26})
    ncols = ["Notification_ID", "Timestamp", "Campaign_ID", "Notification_Type", "To_Name", "To_Email", "Subject", "Channel", "Status"]
    _sheet(wb, "NOTIFICATIONS", ncols, [[r[c] for c in ncols] for r in conn.execute("SELECT * FROM notifications ORDER BY Timestamp")],
           {"Subject": 80, "To_Email": 30, "Timestamp": 20})

    # ---- LISTS (dropdown sources) + data validation on the request sheet
    lists = {"Regions": C.REGIONS, "Campaign_Types": C.CAMPAIGN_TYPES, "Objectives": C.OBJECTIVES, "Channels": C.CHANNELS,
             "Products": C.PRODUCTS, "Approval_Statuses": C.APPROVAL_STATUSES, "Execution_Statuses": C.EXECUTION_STATUSES,
             "Roles": C.ALL_ROLES, "Customer_Segments": C.SEGMENTS}
    ws_li = wb.create_sheet("LISTS")
    for ci, (k, v) in enumerate(lists.items(), start=1):
        ws_li.cell(row=1, column=ci, value=k).font = Font(bold=True)
        ws_li.column_dimensions[get_column_letter(ci)].width = 26
        for ri, item in enumerate(v, start=2):
            ws_li.cell(row=ri, column=ci, value=item)
    ref = lambda key: f"=LISTS!${get_column_letter(list(lists).index(key) + 1)}$2:${get_column_letter(list(lists).index(key) + 1)}${len(lists[key]) + 1}"  # noqa: E731
    for header, key in (("Region", "Regions"), ("Campaign_Type", "Campaign_Types"), ("Objective", "Objectives"),
                        ("Product", "Products"), ("Approval_Status", "Approval_Statuses"), ("Execution_Status", "Execution_Statuses")):
        col = get_column_letter(cols.index(header) + 1)
        dv = DataValidation(type="list", formula1=ref(key), allow_blank=True)
        wb["CAMPAIGN_REQUESTS"].add_data_validation(dv)
        dv.add(f"{col}2:{col}1000")
    dv = DataValidation(type="decimal", operator="greaterThan", formula1="0", allow_blank=True,
                        error="Budget must be greater than zero", showErrorMessage=True)
    wb["CAMPAIGN_REQUESTS"].add_data_validation(dv)
    dv.add("N2:N1000")

    # ---- DASHBOARD_DATA (Looker Studio source): KPI formulas
    wd = wb.create_sheet("DASHBOARD_DATA")
    kp = [
        ("Total Campaigns", f"=COUNTA(CAMPAIGN_REQUESTS!A2:A{n})", "0"),
        ("Pending Approvals", f'=COUNTIF(CAMPAIGN_REQUESTS!V2:V{n},"Pending Approval")', "0"),
        ("Approved Campaigns", f'=COUNTIF(CAMPAIGN_REQUESTS!V2:V{n},"Fully Approved")', "0"),
        ("Campaign Spend", f"=SUM(CAMPAIGN_PERFORMANCE!E2:E{pn})", INR),
        ("Leads Generated", f"=SUM(CAMPAIGN_PERFORMANCE!H2:H{pn})", "#,##0"),
        ("Qualified Leads", f"=SUM(CAMPAIGN_PERFORMANCE!I2:I{pn})", "#,##0"),
        ("Conversions", f"=SUM(CAMPAIGN_PERFORMANCE!J2:J{pn})", "#,##0"),
        ("Revenue", f"=SUM(CAMPAIGN_PERFORMANCE!K2:K{pn})", INR),
        ("Conversion Rate", "=IF(B6>0,B8/B6,\"\")", "0.0%"),                 # Conversions / Leads
        ("Cost per Lead", "=IF(B6>0,B5/B6,\"\")", '"₹"#,##0.00'),            # Spend / Leads
        ("ROAS", "=IF(B5>0,B9/B5,\"\")", '0.00"x"'),                         # Revenue / Spend
        ("ROI", "=IF(B5>0,(B9-B5)/B5,\"\")", "0.0%"),                        # (Revenue - Spend) / Spend
    ]
    wd.append(["KPI", "Value"])
    for name, f, fmt in kp:
        wd.append([name, f])
        wd.cell(row=wd.max_row, column=2).number_format = fmt
    wd.column_dimensions["A"].width = 26; wd.column_dimensions["B"].width = 20
    for c in (1, 2):
        wd.cell(row=1, column=c).font = Font(bold=True, color="FFFFFF"); wd.cell(row=1, column=c).fill = PatternFill("solid", fgColor=NAVY)
    wd["A15"] = "Region"; wd["B15"] = "Campaigns"; wd["C15"] = "Spend"; wd["D15"] = "Revenue"
    for c in "ABCD":
        wd[f"{c}15"].font = Font(bold=True)
    for i, reg in enumerate(C.REGIONS, start=16):
        wd[f"A{i}"] = reg
        wd[f"B{i}"] = f"=COUNTIF(CAMPAIGN_REQUESTS!C2:C{n},A{i})"
        wd[f"C{i}"] = f"=SUMIF(CAMPAIGN_PERFORMANCE!C2:C{pn},A{i},CAMPAIGN_PERFORMANCE!E2:E{pn})"
        wd[f"D{i}"] = f"=SUMIF(CAMPAIGN_PERFORMANCE!C2:C{pn},A{i},CAMPAIGN_PERFORMANCE!K2:K{pn})"
        wd[f"C{i}"].number_format = INR; wd[f"D{i}"].number_format = INR
    wd.column_dimensions["C"].width = 18; wd.column_dimensions["D"].width = 18

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


if __name__ == "__main__":
    from engine import db
    out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Nova_Marketing_Governance_DB.xlsx")
    with open(out, "wb") as f:
        f.write(build(db.connect()))
    print("wrote", out)
